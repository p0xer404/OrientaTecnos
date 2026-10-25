#!/usr/bin/env python3
"""Genera excel/buscador_estudios_cataluna.xlsx a partir de data/*.json (ejecutar después de update_data.py)."""
import json, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
D = lambda n: json.load(open(os.path.join(ROOT, "data", n + ".json"), encoding="utf-8"))
g, fp, us, meta = D("grados"), D("fp"), D("universidades"), D("meta")
ba, ar = D("bachilleratos"), D("artisticos_superiores")
J = lambda v: ", ".join(v) if isinstance(v, list) else ("; ".join(f"{k}: {d['1o']} (1º) / {d['2o']} (2º)" for k, d in v.items()) if isinstance(v, dict) else v)
for x in ba + ar:
    for k, v in list(x.items()): x[k] = J(v)
F = "Arial"; H = Font(name=F, bold=True, color="FFFFFF"); HF = PatternFill("solid", fgColor="1F4E78"); B = Font(name=F)
wb = Workbook(); ws = wb.active; ws.title = "Léeme"
rows = [["Buscador de estudios – Cataluña " + meta["curso"]], [], ["Generado", meta["generado"]], ["Grados", meta["num_grados"]], ["Bachilleratos (área metropolitana)", meta["num_bachilleratos"]], ["Centros artísticos superiores", meta["num_artisticos"]],
        ["Grados con precio", meta["grados_con_precio"]], ["Ofertas FP", meta["num_fp"]], [], ["Fuentes"]]
rows += [[k, v] for k, v in meta["fuentes"].items()]
rows += [[], ["Notas"], ["", "Nota de corte 5 = entraron todos los que lo pidieron."],
         ["", "Precio = primer curso (60 créditos). En públicas no incluye tasas (~140 €). Mira 'Curso precio' y 'Incluye tasas'."],
         ["", "Las privadas (Esade, IQS, La Salle, Blanquerna, UIC, Abat Oliba) tienen admisión propia: no tienen nota de corte."],
         ["", "Algunos privados ya publican el curso 2027-28: se indica en 'Curso precio'."], [], ["Centros sin precio localizado"]]
rows += [["", c] for c in meta["centros_sin_precio"]]
for r in rows: ws.append(r)
for row in ws.iter_rows():
    for c in row: c.font = B; c.alignment = Alignment(wrap_text=True, vertical="top")
ws["A1"].font = Font(name=F, bold=True, size=14); ws.column_dimensions["A"].width = 28; ws.column_dimensions["B"].width = 110

def sheet(title, data, cols, widths, numfmt={}):
    w = wb.create_sheet(title); w.append([c[1] for c in cols])
    for x in data: w.append([x.get(c[0]) for c in cols])
    for j, c in enumerate(w[1], 1):
        c.font = H; c.fill = HF; c.alignment = Alignment(wrap_text=True, vertical="center"); w.column_dimensions[get_column_letter(j)].width = widths[j - 1]
    for row in w.iter_rows(min_row=2):
        for j, c in enumerate(row):
            c.font = B
            if cols[j][0] in numfmt: c.number_format = numfmt[cols[j][0]]
    w.freeze_panes = "C2"; w.auto_filter.ref = w.dimensions

sheet("Grados", g, [("grado","Grado"),("universidad","Univ."),("centro","Centro"),("tipo_centro","Tipo"),("area","Área"),("municipio","Municipio"),("plazas","Plazas"),
      ("nota_corte_2026","Nota 2026"),("nota_corte_2025","Nota 2025"),("nota_corte_2024","Nota 2024"),("nota_corte_2023","Nota 2023"),
      ("precio_1r_curso","Precio 1r curso (€)"),("precio_credito","€/crédito"),("precio_curso_ref","Curso precio"),("precio_incluye_tasas","Incluye tasas"),
      ("admision","Admisión"),("duracion","Duración"),("idiomas","Idiomas"),("nota","Observaciones"),("precio_fuente","Fuente precio"),("url","Web")],
      [60,8,22,22,26,20,8,9,9,9,9,13,10,10,9,30,14,18,40,40,40],
      {"nota_corte_2026":"0.000","nota_corte_2025":"0.000","nota_corte_2024":"0.000","nota_corte_2023":"0.000","precio_1r_curso":"#,##0 €","precio_credito":"#,##0.00 €"})
sheet("FP", fp, [("ciclo","Ciclo"),("nivel","Nivel"),("familia","Familia"),("horas","Horas"),("centro","Centro"),("titularidad","Titularidad"),("municipio","Municipio"),
      ("comarca","Comarca"),("modalidad","Modalidad"),("precio_curso","Precio/curso (€)"),("precio_nota","Nota precio"),("telefono","Teléfono"),("email","Email"),
      ("web","Web"),("direccion","Dirección"),("lat","Lat"),("lon","Lon"),("preinscripcion","Preinscripción")],
      [50,20,28,7,34,11,20,18,11,12,22,12,26,28,30,9,9,30], {"precio_curso":"#,##0 €"})
sheet("Bachilleratos", ba, [("centro","Centro"),("municipio","Municipio"),("titularidad","Titularidad"),("km_desde_terrassa","Km desde Terrassa"),("modalidades","Modalidades"),("ciencias_tecnologia","Ciencias y Tecnología"),("humanidades_sociales","Humanidades y CC. Sociales"),("artes","Artes"),("general","General"),("alumnos_por_modalidad","Alumnos por modalidad (2025-26)"),
      ("turnos","Turnos"),("plazas_1o","Plazas 1º"),("plazas_preinscripcion_1o","Plazas preinscripción 1º"),("asignados_preinscripcion_1o","Asignados"),("asignados_1a_opcion_1o","Asignados 1ª opción"),
      ("ocupacion_preinscripcion","Ocupación"),("precio_nota","Precio"),("comarca","Comarca"),("direccion","Dirección"),("telefono","Teléfono"),("email","Email"),("web","Web"),("lat","Lat"),("lon","Lon")],
      [34,20,12,9,40,10,10,8,8,55,14,9,11,10,10,10,40,18,30,12,26,30,9,9], {"ocupacion_preinscripcion":"0%"})
sheet("Artísticos superiores", ar, [("centro","Centro"),("municipio","Municipio"),("titularidad","Titularidad"),("tipo_estudios","Tipo de estudios"),("ensenanzas_superiores","Enseñanzas superiores"),
      ("especialidades","Especialidades"),("ciclos_superiores_artisticos","Ciclos superiores artísticos"),("admision","Admisión"),("precio_1r_curso","Precio 1r curso (€)"),
      ("precio_credito","€/crédito"),("precio_ciclos","Precio ciclos (€/curso)"),("precio_nota","Nota precio"),("telefono","Teléfono"),("email","Email"),("web","Web"),("direccion","Dirección")],
      [40,20,14,40,28,40,40,45,12,10,12,45,12,26,28,30], {"precio_1r_curso":"#,##0 €","precio_credito":"#,##0.00 €","precio_ciclos":"#,##0 €"})
sheet("Universidades", us, [("siglas","Siglas"),("nombre","Nombre"),("tipo","Tipo"),("num_grados","Nº grados")], [10,55,10,10])
os.makedirs(os.path.join(ROOT, "excel"), exist_ok=True)
out = os.path.join(ROOT, "excel", "buscador_estudios_cataluna.xlsx"); wb.save(out); print("ok", out)
